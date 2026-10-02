#!/usr/bin/env python3
"""
pptx_botones_helper.py — Botones y menús de navegación nativos en .pptx
usando `shape.click_action.target_slide` de python-pptx.

Por qué esto SÍ es seguro de automatizar (a diferencia de las animaciones de
aparición, ver pptx_dinamizar.py): un "Action Setting" de PowerPoint
(`<p:nvSpPr>/<p:cNvPr>/<a:hlinkClick>`) es un hipervínculo interno normal, el
mismo mecanismo que un enlace de texto — no vive en el árbol `<p:timing>` de
animaciones que es imposible de validar sin PowerPoint real. python-pptx lo
expone de forma nativa y documentada (`click_action.target_slide`), y el
resultado es un `.pptx` que cualquier versión de PowerPoint (no solo
2019+, a diferencia de Morph) sabe interpretar: es funcionalidad presente
desde Office 2007.

Qué SÍ resuelve: menús de navegación tipo "índice clicable" (una diapositiva
con secciones, cada una salta a su parte del deck), botones "volver al menú"
en cada sección, y navegación tipo Prezi por botones en vez de orden lineal —
los mismos patrones vistos en los tutoriales de menús interactivos/acordeón de
Jazz Guzman, mencionados en tecnica-morph-real.md como pendientes de
verificar en detalle por el bloqueo de subtítulos de YouTube. La mecánica de
navegación (Action Settings) es API estable y documentada de python-pptx —no
depende de haber visto esos videos completos para implementarse con
confianza— lo que sí queda pendiente de esos videos es el AGRUPADO VISUAL
(acordeón, iconos animados) más elaborado que se ve en pantalla.

Qué NO resuelve: el efecto visual de "abrir/cerrar" un acordeón con animación
(eso sí requiere `<p:timing>` con disparadores, fuera del alcance seguro de
esta skill) — aquí el clic cambia de diapositiva al instante, sin transición
de apertura. Para simular una sensación de menú desplegable sin animación de
apertura, combinar con una transición suave (`fade` o `push`) en
pptx_dinamizar.py sobre la diapositiva de destino.

Uso:
    from pptx import Presentation
    from pptx_botones_helper import agregar_boton, agregar_boton_inicio, construir_menu_interactivo

    prs = Presentation("entrada.pptx")

    # Menú en la diapositiva 0, con 3 secciones que saltan a las diapositivas 1, 4 y 7
    construir_menu_interactivo(
        prs, indice_menu=0,
        items=[("Definición y epidemiología", 1), ("Diagnóstico", 4), ("Tratamiento", 7)],
        agregar_boton_regreso=True,
    )
    prs.save("con_menu.pptx")
"""

from __future__ import annotations

from dataclasses import dataclass

try:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "Este módulo requiere python-pptx (pip install python-pptx --break-system-packages)"
    ) from exc


@dataclass
class EstiloBoton:
    color_fondo: RGBColor = RGBColor(0x16, 0x27, 0x3F)
    color_texto: RGBColor = RGBColor(0xFF, 0xFF, 0xFF)
    color_borde: RGBColor = RGBColor(0x3A, 0xA1, 0xFF)
    tamano_fuente: int = 16
    forma: int = MSO_SHAPE.ROUNDED_RECTANGLE


ESTILO_DEFECTO = EstiloBoton()


def agregar_boton(
    slide,
    x_in: float,
    y_in: float,
    w_in: float,
    h_in: float,
    texto: str,
    target_slide,
    estilo: EstiloBoton = ESTILO_DEFECTO,
    nombre: str | None = None,
):
    """Añade una forma con texto en `slide` que al hacer clic salta a
    `target_slide` (un objeto Slide, ej. `prs.slides[4]`). Devuelve la forma
    creada por si se necesita seguir personalizándola."""
    shape = slide.shapes.add_shape(estilo.forma, Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in))
    shape.fill.solid()
    shape.fill.fore_color.rgb = estilo.color_fondo
    shape.line.color.rgb = estilo.color_borde
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.text = texto
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(estilo.tamano_fuente)
    p.font.bold = True
    p.font.color.rgb = estilo.color_texto

    # El propio mecanismo de hipervínculo interno de PowerPoint — API pública
    # y documentada de python-pptx, no XML hecho a mano.
    shape.click_action.target_slide = target_slide

    if nombre:
        shape.name = nombre
    return shape


def agregar_boton_inicio(
    slide,
    target_slide,
    tamano_in: float = 0.5,
    margen_in: float = 0.3,
    estilo: EstiloBoton | None = None,
):
    """Añade un pequeño botón circular "inicio" (⌂) en la esquina inferior
    izquierda de `slide` que salta de vuelta a `target_slide` — el patrón
    "volver al menú" de las infografías/menús interactivos. Pensado para
    llamarse una vez por cada diapositiva de contenido del deck."""
    estilo = estilo or ESTILO_DEFECTO
    slide_h_in = slide.part.package.presentation_part.presentation.slide_height / 914400
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(margen_in),
        Inches(slide_h_in - tamano_in - margen_in),
        Inches(tamano_in),
        Inches(tamano_in),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = estilo.color_borde
    shape.line.fill.background()
    tf = shape.text_frame
    tf.text = "⌂"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = estilo.color_texto
    shape.click_action.target_slide = target_slide
    shape.name = "BotonInicio"
    return shape


def construir_menu_interactivo(
    prs,
    indice_menu: int,
    items: list[tuple[str, int]],
    agregar_boton_regreso: bool = True,
    columna_x_in: float = 1.0,
    ancho_boton_in: float = 8.0,
    alto_boton_in: float = 0.9,
    y_inicial_in: float = 2.0,
    espacio_in: float = 0.25,
    estilo: EstiloBoton = ESTILO_DEFECTO,
):
    """Construye un menú de navegación completo: un botón por cada
    `(texto, indice_diapositiva_destino)` en `items`, apilados verticalmente
    en la diapositiva `indice_menu`. Si `agregar_boton_regreso` es True,
    añade además un botón "⌂ volver al menú" en cada diapositiva destino.

    No decide el texto ni el orden de las secciones — eso lo define quien
    llama (normalmente reflejando el índice/tabla de contenido que
    `presentaciones-cientificas` ya redactó), esta función solo construye la
    navegación clicable de forma segura."""
    slide_menu = prs.slides[indice_menu]
    y = y_inicial_in
    for i, (texto, idx_destino) in enumerate(items):
        destino = prs.slides[idx_destino]
        agregar_boton(
            slide_menu, columna_x_in, y, ancho_boton_in, alto_boton_in,
            texto, destino, estilo=estilo, nombre=f"BotonMenu_{i}",
        )
        y += alto_boton_in + espacio_in

        if agregar_boton_regreso:
            agregar_boton_inicio(destino, slide_menu, estilo=estilo)
