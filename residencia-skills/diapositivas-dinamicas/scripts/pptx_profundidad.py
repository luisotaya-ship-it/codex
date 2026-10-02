#!/usr/bin/env python3
"""
pptx_profundidad.py — Añade sombra suave (profundidad tipo tarjeta flotante) a
las formas rellenas de un .pptx (roundRect / rect usadas como tarjetas, cajas
de algoritmo, cuadrantes) sin tocar texto, imágenes ni marcadores de título.

Es la parte "3D barato pero seguro" del lado .pptx: un <a:outerShdw> en
spPr/effectLst es un elemento DrawingML único, extremadamente estándar (lo usa
cualquier plantilla de PowerPoint con "Estilos rápidos de forma"), muy distinto
en riesgo de un árbol de animación <p:timing> — por eso esto sí se hace por XML
directo con confianza, y las animaciones de aparición no (ver pptx_dinamizar.py).

Uso:
    python3 pptx_profundidad.py entrada.pptx salida.pptx
    python3 pptx_profundidad.py entrada.pptx salida.pptx --intensidad fuerte
"""

from __future__ import annotations

import argparse
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NSMAP = {"a": A_NS}

INTENSIDADES = {
    # blurRad y dist en EMU (914400 EMU = 1 in). dir en 1/60000 de grado.
    "sutil": {"blur_pt": 10, "dist_pt": 2, "alpha": 28000},
    "media": {"blur_pt": 16, "dist_pt": 4, "alpha": 35000},
    "fuerte": {"blur_pt": 22, "dist_pt": 6, "alpha": 45000},
}

# Formas que tiene sentido "levantar" con sombra: cajas/tarjetas/rects con
# relleno propio. Se excluyen marcadores de texto (título, cuerpo, pie) porque
# ahí una sombra solo añade ruido visual, y se excluye el fondo de página
# completa (evita sombrear la plantilla institucional).
SHAPE_TYPES_ELEGIBLES = {
    MSO_SHAPE_TYPE.AUTO_SHAPE,
    MSO_SHAPE_TYPE.FREEFORM,
}


def _pt_to_emu(pt: float) -> int:
    return int(pt * 12700)


def _has_own_fill(shape) -> bool:
    try:
        return shape.fill.type is not None and shape.fill.type != 5  # 5 = BACKGROUND
    except Exception:
        return False


def _is_full_slide_background(shape, slide_width: int, slide_height: int) -> bool:
    try:
        return (
            shape.left == 0
            and shape.top == 0
            and abs(shape.width - slide_width) < _pt_to_emu(2)
            and abs(shape.height - slide_height) < _pt_to_emu(2)
        )
    except Exception:
        return False


def _add_shadow(shape, blur_pt: float, dist_pt: float, alpha: int) -> bool:
    sp_pr = shape._element.spPr
    if sp_pr is None:
        return False
    existing = sp_pr.find(f"{{{A_NS}}}effectLst")
    if existing is not None:
        return False  # no pisar un efecto que ya trae la plantilla/deck

    effect_lst = etree.SubElement(sp_pr, f"{{{A_NS}}}effectLst")
    shadow = etree.SubElement(effect_lst, f"{{{A_NS}}}outerShdw")
    shadow.set("blurRad", str(_pt_to_emu(blur_pt)))
    shadow.set("dist", str(_pt_to_emu(dist_pt)))
    shadow.set("dir", "5400000")  # hacia abajo
    shadow.set("rotWithShape", "0")
    color = etree.SubElement(shadow, f"{{{A_NS}}}srgbClr")
    color.set("val", "000000")
    alpha_el = etree.SubElement(color, f"{{{A_NS}}}alpha")
    alpha_el.set("val", str(alpha))
    return True


def aplicar_profundidad(input_pptx: Path, output_pptx: Path, intensidad: str = "media") -> int:
    if intensidad not in INTENSIDADES:
        raise SystemExit(f"Intensidad desconocida: {intensidad!r}. Usa: {', '.join(INTENSIDADES)}")
    cfg = INTENSIDADES[intensidad]
    prs = Presentation(str(input_pptx))
    sw, sh = prs.slide_width, prs.slide_height

    tocadas = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type not in SHAPE_TYPES_ELEGIBLES:
                continue
            if _is_full_slide_background(shape, sw, sh):
                continue
            if not _has_own_fill(shape):
                continue
            if _add_shadow(shape, cfg["blur_pt"], cfg["dist_pt"], cfg["alpha"]):
                tocadas += 1

    prs.save(str(output_pptx))
    return tocadas


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("entrada", type=Path)
    parser.add_argument("salida", type=Path)
    parser.add_argument("--intensidad", choices=list(INTENSIDADES), default="media")
    args = parser.parse_args()
    n = aplicar_profundidad(args.entrada, args.salida, args.intensidad)
    print(f"Sombra de profundidad aplicada a {n} forma(s) -> {args.salida}")


if __name__ == "__main__":
    main()
