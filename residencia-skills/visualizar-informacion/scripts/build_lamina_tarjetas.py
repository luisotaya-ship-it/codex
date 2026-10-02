#!/usr/bin/env python3
"""Construye una lamina de N elementos paralelos (3-6 tarjetas en una fila).

Aplica el presupuesto de carga cognitiva de la skill `visualizar-informacion`:
encabezado de color, icono, <=3 vinetas de <=15 palabras, chip de grado de
evidencia, banda de cierre y cita al pie, todo dentro de la zona segura de la
plantilla UNINAVARRA.

Uso:
    python3 build_lamina_tarjetas.py spec.json salida.pptx
    python3 build_lamina_tarjetas.py --esquema
"""

import json
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Geometria (pulgadas) verificada contra la plantilla oficial 2025 ---
LIENZO_W, LIENZO_H = 13.333, 7.5
MARGEN_X = 0.55
ANCHO_UTIL = LIENZO_W - 2 * MARGEN_X
TITULO_TOP, TITULO_H = 1.55, 0.75
FILA_TOP = 2.40
CIERRE_H = 0.62
CITA_TOP = 7.05          # convencion oficial 2025; ver SKILL.md si el deck usa 6.45
FONDO_FILA = 6.85        # limite inferior de las tarjetas cuando NO hay banda de cierre
GAP = 0.18

NAVY = RGBColor(0x0E, 0x28, 0x41)
GRIS_TEXTO = RGBColor(0x3A, 0x3A, 0x3A)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

LIMITE_PALABRAS = 35
LIMITE_VINETAS = 3
LIMITE_PALABRAS_VINETA = 15
LIMITE_TARJETAS = 6

ESQUEMA = {
    "titulo": "Medidas no farmacologicas: el pilar que sostiene todo lo demas",
    "fondo": "gran sesion/Plantilla_UNINAVARRA_fondo_2025.jpg",
    "cita": "ADA Prof. Practice Committee. Sec. 5. Diabetes Care. 2026;49(Suppl 1):S89-S131.",
    "notas": "Guion del orador para esta lamina.",
    "tarjetas": [
        {
            "titulo": "Actividad fisica",
            "color": "#1D7A3E",
            "icono": "iconos/tenis.png",
            "vinetas": [
                {"texto": ">=150 min/sem aerobico moderado-vigoroso, en >=3 dias",
                 "grado": "5.36 - B"},
                {"texto": "Sin mas de 2 dias consecutivos sin actividad",
                 "grado": "5.36 - B"},
                {"texto": "Interrumpir el sedentarismo cada 30 min", "grado": "5.34 - C"}
            ]
        }
    ],
    "cierre": {
        "mensaje": "El tratamiento no farmacologico es el primer pilar del manejo integral.",
        "puntos": ["Mejor control glucemico", "Menor riesgo cardiovascular",
                   "Mayor calidad de vida"]
    }
}


def hex_a_rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def tinte(h, factor=0.90):
    """Aclara un color hacia el blanco. factor 0.90 = 90% blanco."""
    h = h.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    mez = lambda c: int(c + (255 - c) * factor)
    return RGBColor(mez(r), mez(g), mez(b))


def sin_borde(shape):
    shape.line.fill.background()
    shape.shadow.inherit = False


def caja(slide, x, y, w, h, forma=MSO_SHAPE.ROUNDED_RECTANGLE, relleno=None):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if relleno is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = relleno
    sin_borde(s)
    try:
        s.adjustments[0] = 0.08
    except (IndexError, KeyError):
        pass
    return s


def texto(slide, x, y, w, h, contenido, tam=11, color=GRIS_TEXTO, negrita=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(45720)
    tf.margin_top = tf.margin_bottom = Emu(18288)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = contenido
    r.font.size = Pt(tam)
    r.font.bold = negrita
    r.font.color.rgb = color
    r.font.name = "Arial"
    return tb


def validar(spec):
    avisos = []
    tarjetas = spec.get("tarjetas", [])
    if len(tarjetas) > LIMITE_TARJETAS:
        avisos.append(
            f"{len(tarjetas)} tarjetas en una fila (limite {LIMITE_TARJETAS}). "
            "Divide en dos laminas: mas de 6 columnas deja cada una ilegible.")
    if len(tarjetas) < 2:
        avisos.append("Menos de 2 tarjetas: este patron es para elementos paralelos; "
                      "una sola idea rinde mas como cifra grande o como prosa.")
    for t in tarjetas:
        vin = t.get("vinetas", [])
        palabras = sum(len(v["texto"].split()) for v in vin)
        if palabras > LIMITE_PALABRAS:
            avisos.append(f"'{t['titulo']}': {palabras} palabras (limite {LIMITE_PALABRAS}). "
                          "Recorta adjetivos, nunca criterios ni umbrales.")
        if len(vin) > LIMITE_VINETAS:
            avisos.append(f"'{t['titulo']}': {len(vin)} vinetas (limite {LIMITE_VINETAS}).")
        for v in vin:
            n = len(v["texto"].split())
            if n > LIMITE_PALABRAS_VINETA:
                avisos.append(f"'{t['titulo']}': una vineta tiene {n} palabras "
                              f"(limite {LIMITE_PALABRAS_VINETA}).")
        if not any(v.get("grado") for v in vin):
            avisos.append(f"'{t['titulo']}': ninguna vineta lleva grado/codigo de evidencia. "
                          "Si la fuente lo trae, no se sacrifica por estetica.")
    if not spec.get("cita"):
        avisos.append("Sin cita al pie: pon la fuente real o reporta que esta pendiente.")
    return avisos


def construir(spec, salida):
    pres = Presentation()
    pres.slide_width = Inches(LIENZO_W)
    pres.slide_height = Inches(LIENZO_H)
    slide = pres.slides.add_slide(pres.slide_layouts[6])

    fondo = spec.get("fondo")
    if fondo:
        try:
            slide.shapes.add_picture(fondo, 0, 0,
                                     Inches(LIENZO_W), Inches(LIENZO_H))
        except (FileNotFoundError, OSError):
            print(f"AVISO: no se encontro el fondo '{fondo}'. "
                  "La lamina sale sin plantilla institucional.")

    texto(slide, MARGEN_X, TITULO_TOP, ANCHO_UTIL, TITULO_H,
          spec["titulo"], tam=28, color=NAVY, negrita=True, anchor=MSO_ANCHOR.MIDDLE)

    tarjetas = spec["tarjetas"]
    n = len(tarjetas)
    hay_cierre = bool(spec.get("cierre"))
    tope_bottom = (CITA_TOP - 0.20 - CIERRE_H - 0.15) if hay_cierre else FONDO_FILA
    ancho = (ANCHO_UTIL - GAP * (n - 1)) / n
    enc_h = 0.52

    # Auto-ajuste: la tarjeta termina donde termina su contenido. Estirarla hasta
    # el fondo deja espacio muerto, que es justo lo que hace ver plana la lamina.
    hay_icono = any(t.get("icono") for t in tarjetas)
    alto_icono = min(1.05, (tope_bottom - FILA_TOP) * 0.32) if hay_icono else 0
    chars_linea = max(14, int((ancho - 0.35) * 96 / 6.2))
    lineas = 0
    for t in tarjetas:
        acum = 0
        for v in t.get("vinetas", []):
            largo = len(v["texto"]) + (len(v.get("grado", "")) + 4 if v.get("grado") else 0)
            acum += max(1, -(-largo // chars_linea))
        lineas = max(lineas, acum)
    alto_texto = lineas * 0.175 + len(max(tarjetas, key=lambda t: len(t.get("vinetas", [])))
                                      .get("vinetas", [])) * 0.09
    fila_h = enc_h + 0.12 + (alto_icono + 0.15 if hay_icono else 0) + alto_texto + 0.18
    fila_h = max(1.8, min(fila_h, tope_bottom - FILA_TOP))
    # Si sobra aire, centra la fila entre el titulo y la banda de cierre en vez
    # de dejar todo el hueco abajo.
    fila_top = FILA_TOP + max(0.0, (tope_bottom - FILA_TOP - fila_h) / 2)
    fila_bottom = fila_top + fila_h
    for i, t in enumerate(tarjetas):
        x = MARGEN_X + i * (ancho + GAP)
        color = t.get("color", "#185FA5")
        rgb = hex_a_rgb(color)

        caja(slide, x, fila_top + enc_h - 0.12, ancho, fila_h - enc_h + 0.12,
             relleno=tinte(color))
        caja(slide, x, fila_top, ancho, enc_h, relleno=rgb)
        texto(slide, x, fila_top, ancho, enc_h, t["titulo"], tam=12,
              color=BLANCO, negrita=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)

        cursor = fila_top + enc_h + 0.12
        icono = t.get("icono")
        if icono:
            try:
                pic = slide.shapes.add_picture(icono, Inches(x), Inches(cursor),
                                               height=Inches(alto_icono))
                pic.left = Inches(x + (ancho - pic.width / 914400) / 2)
                cursor += alto_icono + 0.15
            except (FileNotFoundError, OSError):
                print(f"AVISO: icono no encontrado para '{t['titulo']}': {icono}. "
                      "La tarjeta pierde la doble codificacion (texto + imagen).")

        tb = slide.shapes.add_textbox(Inches(x + 0.10), Inches(cursor),
                                      Inches(ancho - 0.20),
                                      Inches(fila_bottom - cursor - 0.10))
        tf = tb.text_frame
        tf.word_wrap = True
        for j, v in enumerate(t.get("vinetas", [])):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            r = p.add_run()
            r.text = "\u2022  " + v["texto"]
            r.font.size = Pt(10.5)
            r.font.color.rgb = GRIS_TEXTO
            r.font.name = "Arial"
            if v.get("grado"):
                g = p.add_run()
                g.text = "  (" + v["grado"] + ")"
                g.font.size = Pt(8.5)
                g.font.bold = True
                g.font.color.rgb = rgb
                g.font.name = "Arial"

    if hay_cierre:
        c = spec["cierre"]
        top = CITA_TOP - 0.20 - CIERRE_H
        caja(slide, MARGEN_X, top, ANCHO_UTIL, CIERRE_H,
             relleno=RGBColor(0xF1, 0xEF, 0xE8))
        texto(slide, MARGEN_X + 0.15, top, ANCHO_UTIL * 0.52, CIERRE_H,
              "MENSAJE CLAVE  \u2014  " + c["mensaje"], tam=11.5, color=NAVY,
              negrita=True, anchor=MSO_ANCHOR.MIDDLE)
        if c.get("puntos"):
            texto(slide, MARGEN_X + ANCHO_UTIL * 0.55, top, ANCHO_UTIL * 0.44, CIERRE_H,
                  "   \u00b7   ".join(c["puntos"]), tam=10, color=GRIS_TEXTO,
                  align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    if spec.get("cita"):
        texto(slide, 1.30, CITA_TOP, LIENZO_W - 2.60, 0.35, spec["cita"],
              tam=8, color=GRIS_TEXTO, align=PP_ALIGN.CENTER)

    if spec.get("notas"):
        slide.notes_slide.notes_text_frame.text = spec["notas"]

    pres.save(salida)
    return salida


def main():
    if "--esquema" in sys.argv:
        print(json.dumps(ESQUEMA, indent=2, ensure_ascii=False))
        return
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        spec = json.load(f)
    avisos = validar(spec)
    for a in avisos:
        print("AVISO: " + a)
    construir(spec, sys.argv[2])
    print("OK -> " + sys.argv[2])
    if avisos:
        print("Revisa los avisos antes de entregar: son las reglas que separan "
              "una lamina legible de un muro de texto con colores.")


if __name__ == "__main__":
    main()
